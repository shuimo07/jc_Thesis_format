# -*- coding: utf-8 -*-
"""
build_template.py — 生成一份只含格式、不含内容的《成都锦城学院本科毕业论文》模板
================================================================================

产出的 论文模板.docx 覆盖了学院要求的全部结构性页面与版式，学生可以直接在
上面替换文字，不用再自己调格式。

结构：
    封面（校名 / 题目 / 二级学院 / 专业 / 姓名 / 学号 / 指导教师）
    学术诚信声明                页码 I
    版权使用授权书              页码 II
    摘要                       页码 III
    Abstract                   页码 IV
    目录（TOC 域，打开自动刷新） 页码 V
    —— 分节，页码重置为阿拉伯数字 ——
    第一章 绪论                 页码 1
       1.1 / 1.1.1 示例层级
       示例正文、示例图、示例三线表、示例题注
    结论 / 参考文献 / 致谢

用法：
    python build_template.py                 # 输出 ./论文模板.docx
    python build_template.py -o 别的名字.docx
"""

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from jc_thesis_format import (
    ALIGN_BOTH, ALIGN_CENTER, ALIGN_LEFT, IDEO_SPACE, SPEC,
    add_chapter_body_styles, add_toc_styles, build_authorization,
    build_cn_abstract, build_cover, build_en_abstract,
    build_integrity_statement, caption_paragraph, enable_update_fields,
    make_three_line_table, new_p, new_section, apply_page, pf, rf, toc_field,
    toc_heading,
)


def main(out_path):
    doc = Document()
    apply_page(doc.sections[0])
    add_chapter_body_styles(doc)
    add_toc_styles(doc)

    # ---------------------------------------------------------- 封面（第 1 节，无页码）
    build_cover(doc,
                title='【论文题目】——【副标题】',
                college='计算机与软件学院',
                major='软件工程',
                name='【姓名】',
                sid='【学号】',
                advisor='【指导教师】')

    # ------------------------------------------- 第 2 节：前置部分，大写罗马数字页码
    new_section(doc, 'upperRoman', 1)
    build_integrity_statement(doc)
    build_authorization(doc)
    build_cn_abstract(doc, '【此处填写中文摘要，约 300–500 字。】',
                      ('关键词1', '关键词2', '关键词3', '关键词4'))
    build_en_abstract(doc, '[English abstract, about 200-250 words.]',
                      ('keyword 1', 'keyword 2', 'keyword 3', 'keyword 4'))

    # 目录
    toc_heading(doc, '目录')
    toc_field(doc, [(1, '第一章' + IDEO_SPACE + '绪论'),
                    (2, '1.1' + IDEO_SPACE + '研究背景与意义'),
                    (3, '1.1.1' + IDEO_SPACE + '三级标题示例'),
                    (1, '结论'),
                    (1, '参考文献'),
                    (1, '致谢')])

    # ------------------------------------------- 第 3 节：正文，阿拉伯数字页码从 1 起
    new_section(doc, 'decimal', 1)

    def h1(text, page_break=True):
        p = doc.add_paragraph(style='Heading 1')
        pf(p, align=ALIGN_LEFT, before=SPEC['h1_before'], after=0,
           line=SPEC['h1_line'], page_break=page_break)
        rf(p.add_run(text), cjk='黑体', size=SPEC['h1_size'], bold=True)
        return p

    def h2(text):
        p = doc.add_paragraph(style='Heading 2')
        pf(p, align=ALIGN_LEFT, before=SPEC['h2_before'], after=0, line=SPEC['body_line'])
        rf(p.add_run(text), cjk='黑体', size=SPEC['h2_size'], bold=True)
        return p

    def h3(text):
        p = doc.add_paragraph(style='Heading 3')
        pf(p, align=ALIGN_LEFT, before=SPEC['h3_before'], after=0, line=SPEC['body_line'])
        rf(p.add_run(text), cjk='黑体', size=SPEC['h3_size'], bold=True)
        return p

    def body(text):
        p = doc.add_paragraph(style='论文正文')
        pf(p, align=ALIGN_BOTH, before=0, after=0, line=SPEC['body_line'],
           first=SPEC['body_first_indent_chars'])
        rf(p.add_run(text), cjk='宋体', size=SPEC['body_size'])
        return p

    h1('第一章' + IDEO_SPACE + '绪论', page_break=False)
    body('这是正文示例段落。宋体小四（12pt），首行缩进 2 字符，两端对齐，'
         '固定行距 22 磅。西文与数字用 Times New Roman，中文用宋体，'
         '替换文字时保持该样式即可，不要把整段文字变成另一种字体。')
    h2('1.1' + IDEO_SPACE + '研究背景与意义')
    body('二级标题为黑体四号（14pt）加粗、左对齐、段前 5 磅。')
    h3('1.1.1' + IDEO_SPACE + '三级标题示例')
    body('三级标题为黑体小四（12pt）加粗、左对齐、段前 7 磅。')

    # 示例图（无图片，用文字框位代替，学生替换成自己的图即可）
    p = new_p(doc, align=ALIGN_CENTER, before=6, after=6)
    rf(p.add_run('［此处插入图片，宽度统一 15.0 cm］'), cjk='宋体', size=12)
    caption_paragraph(doc, '图 1-1' + IDEO_SPACE + '图片题注示例（图题在图下方）', kind='fig')

    # 示例三线表
    caption_paragraph(doc, '表 1-1' + IDEO_SPACE + '表格题注示例（表题在表上方）', kind='tab')
    make_three_line_table(doc, [
        ['字段', '含义', '示例值'],
        ['user_id', '用户唯一标识', '10023841'],
        ['platform', '内容平台', 'bilibili'],
        ['sentiment', '情感极性', '正面 / 中性 / 负面'],
    ])
    new_p(doc, before=6, after=6)

    body('表格为三线表：顶线与底线 1.5 磅、表头下 0.75 磅，无竖线与内部横线，'
         '表内文字宋体五号居中。')

    h1('结论')
    body('【此处填写结论】')
    h1('参考文献')
    for i, txt in enumerate([
        '[1] 作者. 文献题名[J]. 期刊名, 2024, 12(3): 45-52.',
        '[2] 作者. 书名[M]. 北京: 出版社, 2023: 100-115.',
        '[3] Author A, Author B. Title of the paper[C]. Proceedings, 2024: 1-8.',
    ], 1):
        p = doc.add_paragraph(style='参考文献')
        pf(p, align=ALIGN_BOTH, before=0, after=0, line=SPEC['body_line'])
        rf(p.add_run(txt), cjk='宋体', size=SPEC['body_size'])
    h1('致谢')
    body('【此处填写致谢】')

    enable_update_fields(doc)
    doc.save(out_path)
    print('saved', out_path)
    print('paras', len(doc.paragraphs), 'tables', len(doc.tables),
          'sections', len(doc.sections))


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('-o', '--out', default='论文模板.docx')
    a = ap.parse_args()
    main(a.out)
